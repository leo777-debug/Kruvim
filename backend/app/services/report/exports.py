"""Portable branded report downloads with embedded, measured charts. No model calls."""
from __future__ import annotations

import base64
import io
import re
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw, ImageFont


def font_path():
    return next((str(p) for p in [Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                                 Path("C:/Windows/Fonts/arial.ttf")] if p.exists()), None)


def chart_sets(results):
    score = results.get("score") or {}
    sets = [("Audience score", [(k, float(score[v]), 10) for k, v in [("After discussion", "mean"),
                                ("First impression", "first_impression")] if score.get(v) is not None])]
    rows = [(s.get("label") or f"Segment {i + 1}", float(s.get("retention") or 0) * 100, 100)
            for i, s in enumerate((results.get("heatmap") or {}).get("segments", []))]
    sets.append(("Attention heatmap  Percentage still watching", rows))
    groups = (results.get("groups") or {}).get("region", [])
    sets.append(("Audience segments", [(g.get("label") or g.get("key") or "Segment", float(g.get("score") or 0), 10) for g in groups]))
    return [(title if len(rows) <= 12 else f"{title}  Part {i // 12 + 1}", rows[i:i + 12])
            for title, rows in sets for i in range(0, max(1, len(rows)), 12)]


def chart_png(title, rows, accent="#2155cd"):
    if not re.fullmatch(r"#[0-9a-fA-F]{6}", accent or ""):
        accent = "#2155cd"
    height = 120 + max(1, len(rows)) * 58
    image = Image.new("RGB", (1200, height), "white")
    draw = ImageDraw.Draw(image)
    path = font_path()
    font = ImageFont.truetype(path, 23) if path else ImageFont.load_default(size=23)
    small = ImageFont.truetype(path, 20) if path else ImageFont.load_default(size=20)
    draw.text((25, 22), title, fill="#111827", font=font)
    if not rows:
        draw.text((25, 95), "No measurements available", fill="#6b7280", font=small)
    for i, (label, value, maximum) in enumerate(rows):
        y = 90 + i * 58
        # Bound labels to their column by measured glyph width.
        label = str(label)
        while draw.textlength(label, font=small) > 410:
            label = label[:-4] + "..."
        draw.text((25, y), label, fill="#374151", font=small)
        draw.rounded_rectangle((455, y, 1075, y + 26), radius=4, fill="#eef0f4")
        width = 620 * min(1, max(0, value / maximum))
        if width:
            color = accent if maximum == 10 else (int(220 - value * 1.5), int(100 + value), 100)
            draw.rounded_rectangle((455, y, 455 + width, y + 26), radius=4, fill=color)
        draw.text((1090, y), f"{value:.1f}" + ("%" if maximum == 100 else "/10"), fill="#111827", font=small)
    out = io.BytesIO()
    image.save(out, "PNG")
    return out.getvalue()


def plain_markdown(line):
    line = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", line)
    line = re.sub(r"\[([^\]]+)\]\(https?://[^)]*\)", r"\1", line)
    return line.replace("**", "").replace("`", "").lstrip("> ")


def export_report(markdown: str, results: dict, branding: dict, format_: str, logo: bytes | None = None) -> bytes:
    brand = str(branding.get("product_name") or "Kruvim")[:60]
    footer = str(branding.get("report_footer") or "")[:300]
    charts = [(t, rows, chart_png(t, rows, branding.get("accent"))) for t, rows in chart_sets(results)]
    if format_ == "md":
        text = brand + "\n\n" + markdown + "\n\n## Key charts\n"
        if branding.get("logo_url"):
            text = f"![{brand}]({branding['logo_url']})\n\n" + text
        for title, rows, png in charts:
            text += f"\n### {title}\n\n![{title}](data:image/png;base64,{base64.b64encode(png).decode()})\n\n"
            text += "| Segment | Value | Scale |\n| --- | ---: | ---: |\n"
            text += "".join(f"| {str(label).replace('|', '/')} | {value:.1f} | {maximum} |\n" for label, value, maximum in rows)
        return (text + "\n" + footer).encode("utf-8")
    output = io.BytesIO()
    if format_ == "docx":
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(.7)
        section.left_margin = section.right_margin = Inches(.8)
        doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Arial", Pt(11)
        for style in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
        section.header.paragraphs[0].text = brand
        section.footer.paragraphs[0].text = footer
        if logo:
            doc.add_picture(io.BytesIO(logo), width=Inches(.8))
        for line in markdown.splitlines():
            if line.startswith("# "):
                doc.add_paragraph(plain_markdown(line[2:]), "Title")
            elif line.startswith("## "):
                doc.add_heading(plain_markdown(line[3:]), 1)
            elif line.startswith("### "):
                doc.add_heading(plain_markdown(line[4:]), 2)
            elif line.strip():
                doc.add_paragraph(plain_markdown(line.lstrip("- ")), "List Bullet" if line.startswith("- ") else None)
        doc.add_heading("Key charts", 1)
        for title, _, png in charts:
            paragraph = doc.add_paragraph(title)
            paragraph.paragraph_format.keep_with_next = True
            doc.add_picture(io.BytesIO(png), width=Inches(6.6))
            doc.inline_shapes[-1]._inline.docPr.set("descr", title)
        doc.save(output)
    elif format_ == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Image as PDFImage
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        path = font_path()
        font_name = "KruvimReport"
        if path and font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font_name, path))
        if not path:
            font_name = "Helvetica"
        styles = getSampleStyleSheet()
        for key in ("Normal", "Title", "Heading1", "Heading2"):
            styles[key].fontName = font_name
            styles[key].textColor = colors.black
        styles["Normal"].fontSize, styles["Normal"].leading, styles["Normal"].spaceAfter = 11, 16, 8
        styles["Normal"].splitLongWords = True
        styles.add(styles["Normal"].clone("Footer", fontSize=8, leading=10, spaceAfter=0))
        story = [Paragraph(escape(brand), styles["Normal"]), Spacer(1, 12)]
        if logo:
            story.append(PDFImage(io.BytesIO(logo), width=55, height=55, kind="proportional"))
        for line in markdown.splitlines():
            if not line.strip():
                continue
            style = "Title" if line.startswith("# ") else "Heading1" if line.startswith("## ") else "Heading2" if line.startswith("### ") else "Normal"
            text = plain_markdown(re.sub(r"^#{1,3}\s+", "", line))
            story.append(Paragraph(escape(text), styles[style]))
        story.append(Paragraph("Key charts", styles["Heading1"]))
        for _, _, png in charts:
            with Image.open(io.BytesIO(png)) as im:
                height = im.height * 470 / im.width
            story.append(PDFImage(io.BytesIO(png), width=470, height=height))
            story.append(Spacer(1, 16))
        def page_footer(canvas, doc):
            canvas.saveState()
            canvas.setFont(font_name, 8)
            # Footer wraps within the printable frame instead of running off the page.
            p = Paragraph(escape(footer), styles["Footer"])
            _, h = p.wrap(460, 40)
            p.drawOn(canvas, 55, 28)
            canvas.drawRightString(557, 24, str(doc.page))
            canvas.restoreState()
        SimpleDocTemplate(output, pagesize=letter, leftMargin=55, rightMargin=55, topMargin=50, bottomMargin=65,
                          title=brand + " report", author=brand).build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    else:
        raise ValueError("Unsupported report format")
    return output.getvalue()
