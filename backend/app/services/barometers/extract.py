"""Text out of messy public files, verbatim.

Nothing is cleaned, normalised, deduplicated or "fixed": stray headers, footnotes, merged cells, odd spacing and
mixed languages all reach the agents as published. The only transformations are the ones needed to get text
at all (decoding bytes, reading PDF pages, workbook cells and Word paragraphs).

Files that cannot yield text raise Quarantine with a reason a person can act on."""
from __future__ import annotations

import io
import os
import re
import zipfile
from dataclasses import dataclass, field

MAX_PAGES = 500
MAX_TEXT_CHARS = 3_000_000
MAX_EXPANDED_BYTES = 200 * 1024 * 1024
CHUNK_CHARS = 1500

TEXT_EXT = {".csv", ".tsv", ".txt", ".md", ".json", ".xml", ".html", ".htm"}
ACCEPTED_EXT = TEXT_EXT | {".xlsx", ".pdf", ".docx"}


class Quarantine(Exception):
    """The file was received but no usable text can be read from it."""


@dataclass
class Section:
    text: str
    location: str


@dataclass
class Extracted:
    sections: list[Section]
    encoding: str | None = None
    notes: dict = field(default_factory=dict)

    @property
    def chars(self) -> int:
        return sum(len(s.text) for s in self.sections)


def decode(data: bytes) -> tuple[str, str]:
    """Bytes to text without altering content: UTF-8 (with or without BOM), then UTF-16 with a BOM, then Windows-1252.
    Latin-1 is the last resort because it accepts any byte sequence."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "utf-16"
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc), "utf-8" if enc == "utf-8-sig" else enc
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "latin-1"


def _looks_binary(data: bytes) -> bool:
    sample = data[:8192]
    if not sample:
        return False
    if sample.startswith((b"\xff\xfe", b"\xfe\xff")):
        return False
    controls = sum(1 for b in sample if b < 9 or 13 < b < 32)
    return b"\x00" in sample or controls / len(sample) > 0.05


def _text_file(data: bytes, ext: str) -> Extracted:
    if _looks_binary(data):
        raise Quarantine("This file contains binary data, not text. Export it as CSV, XLSX, PDF or plain text.")
    text, encoding = decode(data)
    if ext in (".html", ".htm"):
        # Markup is not content; everything a reader would see is kept, including navigation and footers.
        text = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", text)
        text = re.sub(r"(?s)<[^>]+>", " ", text)
    return Extracted([Section(text, "")], encoding)


def _xlsx(data: bytes) -> Extracted:
    from openpyxl import load_workbook
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise Quarantine("This workbook expands to more than 200 MB. Split it into smaller files.")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Quarantine:
        raise
    except Exception as exc:
        raise Quarantine("This is not a readable Excel workbook. Re-save it as .xlsx or export it as CSV.") from exc
    sections = []
    try:
        for sheet in book.worksheets:
            lines = []
            for row in sheet.iter_rows(values_only=True):
                lines.append("\t".join("" if v is None else str(v) for v in row).rstrip("\t"))
            if any(lines):
                sections.append(Section("\n".join(lines), f"sheet {sheet.title}"))
    finally:
        book.close()
    return Extracted(sections, notes={"sheets": len(book.sheetnames)})


def _pdf(data: bytes) -> Extracted:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise Quarantine("This PDF is password-protected. Upload an unlocked copy.")
        pages = reader.pages
        sections, empty = [], 0
        for n, page in enumerate(pages[:MAX_PAGES], start=1):
            text = page.extract_text() or ""
            if text.strip():
                sections.append(Section(text, f"page {n}"))
            else:
                empty += 1
    except Quarantine:
        raise
    except Exception as exc:
        raise Quarantine("This PDF is damaged or not a real PDF, so no text could be read from it.") from exc
    notes = {"pages": len(pages), "pages_without_text": empty}
    if len(pages) > MAX_PAGES:
        notes["pages_skipped"] = len(pages) - MAX_PAGES
    if not sections:
        raise Quarantine("No text could be read from this PDF. It is probably scanned images; upload a text PDF or the source data.")
    return Extracted(sections, notes=notes)


def _docx(data: bytes) -> Extracted:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        if archive.getinfo("word/document.xml").file_size > 50 * 1024 * 1024:
            raise Quarantine("This Word document expands to more than 50 MB of text. Split it into smaller files.")
        xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
    except Quarantine:
        raise
    except Exception as exc:
        raise Quarantine("This is not a readable Word document. Re-save it as .docx or export it as PDF.") from exc
    paragraphs = re.findall(r"<w:p[ >].*?</w:p>", xml, re.S)
    text = "\n".join("".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", p)) for p in paragraphs)
    for entity, char in (("&lt;", "<"), ("&gt;", ">"), ("&quot;", '"'), ("&apos;", "'"), ("&amp;", "&")):
        text = text.replace(entity, char)
    return Extracted([Section(text, "")])


def extract(filename: str, data: bytes) -> Extracted:
    ext = os.path.splitext(filename.lower())[1]
    if not data:
        raise Quarantine("The file is empty.")
    if ext in TEXT_EXT:
        out = _text_file(data, ext)
    elif ext == ".xlsx":
        out = _xlsx(data)
    elif ext == ".pdf":
        out = _pdf(data)
    elif ext == ".docx":
        out = _docx(data)
    else:
        raise ValueError(f"Unsupported file type {ext or '(none)'}")
    if not any(s.text.strip() for s in out.sections):
        raise Quarantine("The file contains no readable text.")
    if out.chars > MAX_TEXT_CHARS:
        raise Quarantine("This file holds more than 3 million characters of text. Split it into smaller files.")
    return out


def chunks(extracted: Extracted, size: int = CHUNK_CHARS) -> list[Section]:
    """Split at line boundaries (a long line is cut at `size`) and record where each piece came from."""
    out = []
    for section in extracted.sections:
        lines = section.text.splitlines(keepends=True)
        buf, start, line_no = "", 1, 0
        for line in lines:
            line_no += 1
            while len(line) > size:
                if buf:
                    out.append(Section(buf, _where(section.location, start, line_no - 1)))
                    buf, start = "", line_no
                out.append(Section(line[:size], _where(section.location, line_no, line_no)))
                line = line[size:]
            if buf and len(buf) + len(line) > size:
                out.append(Section(buf, _where(section.location, start, line_no - 1)))
                buf, start = "", line_no
            buf += line
        if buf.strip():
            out.append(Section(buf, _where(section.location, start, line_no)))
    return [c for c in out if c.text.strip()]


def _where(location: str, first: int, last: int) -> str:
    lines = f"line {first}" if first >= last else f"lines {first}-{last}"
    return f"{location}, {lines}" if location else lines
