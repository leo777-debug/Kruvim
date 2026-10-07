"""Native-grain source adapters. Reading a file never fabricates or interpolates missing cells."""
from __future__ import annotations

import calendar
import csv
import io
import ipaddress
import json
import math
import socket
import zipfile
from dataclasses import dataclass
from datetime import UTC, date, datetime
from urllib.parse import urlparse

import httpx

MAX_ROWS = 100_000
MAX_BYTES = 25 * 1024 * 1024


def table(raw: bytes, filename="data.csv", sheet=None):
    if len(raw) > MAX_BYTES:
        raise ValueError("Import exceeds the 25 MB table limit")
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(x.file_size for x in archive.infolist()) > MAX_BYTES * 4:
                raise ValueError("Expanded workbook is too large")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
        try:
            page = workbook[sheet] if sheet else workbook.worksheets[0]
            rows = []
            for cells in page.iter_rows():
                if any(c.data_type == "f" for c in cells):
                    raise ValueError("Export formula cells as values before importing")
                rows.append([v.value.isoformat() if isinstance(v.value, (datetime, date)) else v.value for v in cells])
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError("Import exceeds 100,000 rows")
        finally:
            workbook.close()
    else:
        text = raw.decode("utf-8-sig", errors="strict")
        try:
            dialect = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = []
        for row in csv.reader(io.StringIO(text), dialect):
            rows.append(row)
            if len(rows) > MAX_ROWS + 1:
                raise ValueError("Import exceeds 100,000 rows")
    if not rows:
        raise ValueError("The table is empty")
    header = [str(x or "").strip() for x in rows[0]]
    if not all(header) or len(set(header)) != len(header):
        raise ValueError("Headers must be unique and nonempty")
    body = [row for row in rows[1:] if any(v not in (None, "") for v in row)]
    if any(len(row) > len(header) for row in body):
        raise ValueError("A row has more cells than the header")
    return header, [dict(zip(header, [*row, *([None] * (len(header) - len(row)))], strict=True)) for row in body]


def period(value, mode="date", end=False):
    text = str(value).strip()
    if mode == "year":
        year = int(text)
        return datetime(year, 12 if end else 1, 31 if end else 1, tzinfo=UTC)
    if mode == "month":
        year, month = map(int, text.split("-"))
        return datetime(year, month, calendar.monthrange(year, month)[1] if end else 1, tzinfo=UTC)
    parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    return parsed.astimezone(UTC) if parsed.tzinfo else parsed.replace(tzinfo=UTC)


def parse_table(raw, mapping, filename="data.csv"):
    if not isinstance(mapping.get("dimensions", {}), dict) or any(not isinstance(v, str) for v in mapping.get("dimensions", {}).values()):
        raise ValueError("Dimensions must map names to table column names")
    headers, rows = table(raw, filename, mapping.get("sheet"))
    required = [mapping.get("value_col"), *(mapping.get("dimensions") or {}).values()]
    required += [mapping[k] for k in ("metric_col", "unit_col", "geography_col", "period_col", "period_start_col", "period_end_col") if mapping.get(k)]
    if any(column not in headers for column in required):
        raise ValueError("A mapped column does not exist in the uploaded table")
    if not mapping.get("metric") and not mapping.get("metric_col"):
        raise ValueError("Choose a metric or its column")
    if not mapping.get("period_col") and not (mapping.get("period_start_col") and mapping.get("period_end_col")):
        raise ValueError("Map a native period, or both start and end dates")
    missing = {str(x).strip() for x in mapping.get("missing_values", ["", "..", "NA", "N/A"])}
    observations, skipped = [], 0
    for index, row in enumerate(rows, 2):
        value = row[mapping["value_col"]]
        if value is None or str(value).strip() in missing:
            skipped += 1
            continue
        try:
            numeric = float(value)
            if not math.isfinite(numeric):
                raise ValueError("Non-finite value")
            metric = str(row.get(mapping.get("metric_col"), mapping.get("metric", ""))).strip()
            unit = str(row.get(mapping.get("unit_col"), mapping.get("unit", ""))).strip()
            geography = str(row.get(mapping.get("geography_col"), mapping.get("geography", ""))).strip()
            if not metric or not unit or not geography or len(metric) > 120 or len(unit) > 40 or len(geography) > 80:
                raise ValueError("Metric, unit and geography are required and must fit their limits")
            start = period(row[mapping.get("period_start_col") or mapping["period_col"]], mapping.get("period_format", "date"))
            end = period(row[mapping.get("period_end_col") or mapping["period_col"]], mapping.get("period_format", "date"), end=True)
            if end < start:
                raise ValueError("Period end precedes start")
            dimensions = {key: row[column] for key, column in (mapping.get("dimensions") or {}).items()}
            if any(value in (None, "") for value in dimensions.values()):
                raise ValueError("A dimension cell is missing; it cannot be filled silently")
            observations.append(dict(metric=metric, dimensions=dimensions, value=numeric, unit=unit,
                period_start=start, period_end=end, geography=geography))
        except (KeyError, TypeError, ValueError) as exc:
            raise ValueError(f"Row {index}: {exc}") from exc
    return observations, skipped


@dataclass
class Fetched:
    raw: bytes
    filename: str
    url: str


async def public_url(url):
    import asyncio
    parsed = urlparse(url)
    if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password or parsed.port not in (None, 443):
        raise ValueError("Source feeds require a public HTTPS URL")
    addresses = await asyncio.to_thread(socket.getaddrinfo, parsed.hostname, 443)
    if not addresses or any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
        raise ValueError("Private or local source addresses are not allowed")


class SourceAdapter:
    async def fetch(self, source, client: httpx.AsyncClient) -> Fetched:
        raise NotImplementedError

    def parse(self, file: Fetched, source, mapping=None):
        return parse_table(file.raw, mapping or source.config.get("mapping", {}), file.filename)


class TableAdapter(SourceAdapter):
    async def fetch(self, source, client):
        endpoint = source.config.get("endpoint")
        if not endpoint:
            raise ValueError("Choose the publisher's public file URL and map its columns first")
        await public_url(endpoint)
        # Never follow redirects to internal addresses or silently download an HTML login page.
        async with client.stream("GET", endpoint, follow_redirects=False) as response:
            response.raise_for_status()
            if response.is_redirect:
                raise ValueError("Register the final publisher URL; redirects are not followed")
            chunks, size = [], 0
            async for chunk in response.aiter_bytes():
                size += len(chunk)
                if size > MAX_BYTES:
                    raise ValueError("Source file exceeds 25 MB")
                chunks.append(chunk)
        return Fetched(b"".join(chunks), source.config.get("filename", "source.csv"), endpoint)


class WorldBankAdapter(TableAdapter):
    def parse(self, file, source, mapping=None):
        data = json.loads(file.raw)
        if not isinstance(data, list) or len(data) != 2 or not isinstance(data[1], list):
            raise ValueError("World Bank indicator endpoint did not return observations")
        if int(data[0].get("pages", 1)) != 1:
            raise ValueError("The indicator response is paginated; increase the registered page size before importing")
        rows = []
        for cell in data[1]:
            if cell.get("value") is None:
                continue
            rows.append(dict(metric=source.config["metric"], dimensions={"indicator": cell["indicator"]["id"]},
                value=float(cell["value"]), unit=source.config["unit"], geography=source.country,
                period_start=period(cell["date"], "year"), period_end=period(cell["date"], "year", end=True)))
        return rows, sum(cell.get("value") is None for cell in data[1])


ADAPTERS = {"table": TableAdapter(), "sdmx": TableAdapter(), "world_bank": WorldBankAdapter()}
