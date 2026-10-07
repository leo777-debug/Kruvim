import io
from datetime import timedelta
from zipfile import ZipFile

import pytest
from docx import Document
from pypdf import PdfReader

from app.db.base import utcnow
from app.db.session import session_scope
from app.models import Organization, Report, ResultShare, Simulation

RESULTS = {"score": {"mean": 7.2, "first_impression": 6.4}, "heatmap": {"completion": .65, "segments": [
    {"label": "Opening", "retention": .9}, {"label": "Explanation", "retention": .65}]},
    "groups": {"region": [{"label": "Riyadh", "score": 7.4}, {"label": "Dubai", "score": 6.9}]},
    "models": {"private": "must not leak"}, "provider": {"dry": True}}


@pytest.fixture
async def finished_run(client, auth):
    headers, data = auth
    project = (await client.post("/projects", headers=headers, json={"name": "Reports"})).json()
    run = (await client.post(f"/projects/{project['id']}/simulations", headers=headers,
                            json={"name": "Learning report", "content": {"type": "text", "text": "Python practice"}})).json()
    async with session_scope() as s:
        sim = await s.get(Simulation, run["id"])
        sim.status, sim.report_status, sim.results = "completed", "done", RESULTS
        sim.config = {"private_api_key": "must not leak"}
        org = await s.get(Organization, data["orgs"][0]["id"])
        previous = org.settings
        org.settings = {**(previous or {}), "branding": {"product_name": "Acme Studio", "accent": "#1955aa", "report_footer": "Acme audience research"}}
        s.add(Report(simulation_id=sim.id, status="done", title="Learning report", summary="The audience score was 7.2.",
                     markdown="# Learning report\n\nThe audience score was 7.2.\n\n## Evidence\n\nDaily practice helps. [node:analysis:example]\n"))
    yield run["id"], headers
    async with session_scope() as s:
        org = await s.get(Organization, data["orgs"][0]["id"])
        org.settings = previous


async def test_downloads_branding_and_embedded_charts(client, finished_run):
    sim_id, headers = finished_run
    pdf = await client.get(f"/simulations/{sim_id}/report/download?format=pdf", headers=headers)
    assert pdf.status_code == 200, pdf.text[:200] if pdf.status_code != 200 else ""
    reader = PdfReader(io.BytesIO(pdf.content))
    text = " ".join(p.extract_text() for p in reader.pages)
    assert "Acme Studio" in text and "Daily practice helps" in text and "Acme audience research" in text
    assert "node:" not in text and "edge:" not in text
    assert "Sources" in text and "[1]" in text
    assert sum(len(p.images) for p in reader.pages) == 3
    word = await client.get(f"/simulations/{sim_id}/report/download?format=docx", headers=headers)
    assert word.status_code == 200
    doc = Document(io.BytesIO(word.content))
    assert doc.sections[0].header.paragraphs[0].text == "Acme Studio"
    assert "node:" not in " ".join(p.text for p in doc.paragraphs)
    assert any(r.font.superscript for p in doc.paragraphs for r in p.runs)
    assert len(doc.inline_shapes) == 3
    assert "Learning report" in [p.text for p in doc.paragraphs]
    with ZipFile(io.BytesIO(word.content)) as zip_:
        assert len([n for n in zip_.namelist() if n.startswith("word/media/")]) == 3
    md = await client.get(f"/simulations/{sim_id}/report/download?format=md", headers=headers)
    assert "data:image/png;base64" in md.text and "| Riyadh | 7.4 | 10 |" in md.text
    assert "node:" not in md.text and "edge:" not in md.text
    assert "<sup>[1]</sup>" in md.text and "1. Simulation evidence" in md.text
    assert md.headers["cache-control"] == "no-store"
    assert (await client.get(f"/simulations/{sim_id}/report/download?format=exe", headers=headers)).status_code == 422
    assert (await client.get(f"/simulations/{sim_id}/report/download")).status_code == 401


async def test_revocable_expiring_shares_are_read_only_and_tenant_scoped(client, finished_run):
    sim_id, headers = finished_run
    link = (await client.post(f"/simulations/{sim_id}/shares", headers=headers, json={"days": 7})).json()
    token = link["url"].rsplit("/", 1)[1]
    public = await client.get("/shared-results/" + token)
    assert public.status_code == 200
    assert "node:" not in public.text and "edge:" not in public.text
    assert public.json()["report"]["sections"][0]["sources"] == [{"number": 1, "label": "Simulation evidence"}]
    private = (await client.get(f"/simulations/{sim_id}/report", headers=headers)).json()
    assert "node:analysis:example" in private["markdown"]
    assert private["rendered_sections"][0]["sources"][0]["id"] == "analysis:example"
    exported = (await client.get(f"/simulations/{sim_id}/export", headers=headers)).json()
    assert "node:analysis:example" in exported["report_markdown"]
    assert public.json()["branding"]["product_name"] == "Acme Studio"
    assert "must not leak" not in public.text and "config" not in public.json()
    assert public.headers["cache-control"] == "no-store" and public.headers["referrer-policy"] == "no-referrer"
    assert (await client.get(f"/simulations/{sim_id}/shares", headers=headers)).json()[0]["view_count"] == 1
    stranger = (await client.post("/auth/register", json={"email": "share-other@example.com", "password": "correct-horse-battery",
                                "name": "Stranger", "org_name": "Stranger Studio"})).json()
    other = {"Authorization": "Bearer " + stranger["access_token"]}
    assert (await client.delete(f"/simulations/{sim_id}/shares/{link['id']}", headers=other)).status_code == 404
    assert (await client.get(f"/simulations/{sim_id}/report/download", headers=other)).status_code == 404
    assert (await client.post("/shared-results/" + token, json={"markdown": "changed"})).status_code == 405
    assert (await client.delete(f"/simulations/{sim_id}/shares/{link['id']}", headers=headers)).status_code == 200
    assert (await client.get("/shared-results/" + token)).status_code == 404
    expired = (await client.post(f"/simulations/{sim_id}/shares", headers=headers, json={"days": 1})).json()
    async with session_scope() as s:
        row = await s.get(ResultShare, expired["id"])
        row.expires_at = utcnow() - timedelta(seconds=1)
    assert (await client.get("/shared-results/" + expired["url"].rsplit("/", 1)[1])).status_code == 404
    async with session_scope() as s:
        row = await s.get(ResultShare, expired["id"])
        assert row.view_count == 0
